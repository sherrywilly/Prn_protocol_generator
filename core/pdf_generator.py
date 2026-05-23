from io import BytesIO
from pathlib import Path

from django.http import HttpResponse
from django.conf import settings
from django.template.loader import render_to_string

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    Document = None


def _format_date(value):
    return value.strftime('%d/%m/%Y') if value else '—'


def _resident_photo_uri(resident):
    if not getattr(resident, 'photo', None):
        return ''
    try:
        photo_path = Path(resident.photo.path)
    except (ValueError, NotImplementedError):
        return ''
    if not photo_path.exists():
        return ''
    return photo_path.as_uri()


def _set_cell_border(cell, color='999999', size='8'):
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        cell_properties.append(borders)

    for edge_name in ('top', 'left', 'bottom', 'right'):
        edge = borders.find(qn(f'w:{edge_name}'))
        if edge is None:
            edge = OxmlElement(f'w:{edge_name}')
            borders.append(edge)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), size)
        edge.set(qn('w:color'), color)


def _set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.first_child_found_in('w:shd')
    if shading is None:
        shading = OxmlElement('w:shd')
        cell_properties.append(shading)
    shading.set(qn('w:fill'), fill)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        cell_properties.append(margins)

    for side, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        margin = margins.find(qn(f'w:{side}'))
        if margin is None:
            margin = OxmlElement(f'w:{side}')
            margins.append(margin)
        margin.set(qn('w:w'), str(value))
        margin.set(qn('w:type'), 'dxa')


def _clear_cell(cell):
    cell.text = ''


def _set_paragraph_spacing(paragraph, before=0, after=0, line=1.0, alignment=None):
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def _set_document_defaults(document):
    section = document.sections[0]
    # Force A4 to keep pagination consistent across systems/printers.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    normal_style = document.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    normal_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    normal_style.font.size = Pt(10)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)


def _set_table_layout(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    table_properties = table._tbl.tblPr
    table_indent = table_properties.first_child_found_in('w:tblInd')
    if table_indent is None:
        table_indent = OxmlElement('w:tblInd')
        table_properties.append(table_indent)
    table_indent.set(qn('w:w'), '0')
    table_indent.set(qn('w:type'), 'dxa')

    cell_spacing = table_properties.first_child_found_in('w:tblCellSpacing')
    if cell_spacing is None:
        cell_spacing = OxmlElement('w:tblCellSpacing')
        table_properties.append(cell_spacing)
    cell_spacing.set(qn('w:w'), '0')
    cell_spacing.set(qn('w:type'), 'dxa')


def _set_row_height(row, height_cm, exact=False):
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST
def _build_resident_export_filename(resident, extension):
    safe_resident = resident.name.replace(' ', '_')
    return f"MAR_Resident_Profile_{safe_resident}.{extension}"



def _split_lines(value):
    return [line.strip() for line in str(value or '').splitlines() if line.strip()] or ['—']


def _first_line(value, default='—'):
    lines = _split_lines(value)
    return lines[0] if lines else default


def _clean_text(value, fallback='Not recorded'):
    cleaned = str(value or '').strip()
    return cleaned if cleaned else fallback


def _contains_any(value, *terms):
    text = (value or '').lower()
    return any(term in text for term in terms)


def _add_resident_docx_banner(document, text, fill='E8F1EC', color='1F5C4F', font_size=10, border_color='D5E2DA'):
    table = document.add_table(rows=1, cols=1)
    _set_table_layout(table)
    table.rows[0].cells[0].width = Cm(18)
    cell = table.cell(0, 0)
    _clear_cell(cell)
    _set_cell_border(cell, color=border_color, size='6')
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=70, start=100, bottom=70, end=100)

    paragraph = cell.paragraphs[0]
    _set_paragraph_spacing(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor.from_string(color)

    spacer = document.add_paragraph()
    _set_paragraph_spacing(spacer, after=1)


def _add_resident_docx_photo(document, resident):
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, after=2, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    try:
        photo_path = Path(resident.photo.path) if getattr(resident, 'photo', None) else None
    except (ValueError, NotImplementedError):
        photo_path = None

    if photo_path and photo_path.exists():
        run = paragraph.add_run()
        run.add_picture(str(photo_path), width=Cm(2.7), height=Cm(3.5))
        return

    run = paragraph.add_run('[ No Photo ]')
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x73)


def _add_resident_docx_center_text(document, text, size=10, bold=True, color='2D433E', after=4):
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph, after=after, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def _write_resident_docx_kv(cell, label, value, fill='FFFFFF'):
    _clear_cell(cell)
    _set_cell_border(cell, color='E1E8E2', size='6')
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    paragraph = cell.paragraphs[0]
    _set_paragraph_spacing(paragraph, line=1.0)

    label_run = paragraph.add_run(f'{label}: ')
    label_run.bold = True
    label_run.font.name = 'Arial'
    label_run.font.size = Pt(7.2)
    label_run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x4F)

    value_run = paragraph.add_run(value or 'Not recorded')
    value_run.font.name = 'Arial'
    value_run.font.size = Pt(8.2)
    value_run.font.color.rgb = RGBColor(0x22, 0x35, 0x31)


def _add_resident_docx_info_table(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    _set_table_layout(table)
    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        _set_row_height(row, 0.64, exact=False)
        row.cells[0].width = Cm(8.9)
        row.cells[1].width = Cm(8.9)
        _write_resident_docx_kv(row.cells[0], row_values[0][0], row_values[0][1], fill='FFFFFF')
        _write_resident_docx_kv(row.cells[1], row_values[1][0], row_values[1][1], fill='FFFFFF')

    spacer = document.add_paragraph()
    _set_paragraph_spacing(spacer, after=1)


def _add_resident_docx_alert(document, text):
    table = document.add_table(rows=1, cols=1)
    _set_table_layout(table)
    cell = table.cell(0, 0)
    _clear_cell(cell)
    _set_cell_border(cell, color='E6BABA', size='8')
    _set_cell_shading(cell, 'FFF1F1')
    _set_cell_margins(cell, top=80, start=90, bottom=80, end=90)

    paragraph = cell.paragraphs[0]
    _set_paragraph_spacing(paragraph, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(0xB8, 0x00, 0x00)

    spacer = document.add_paragraph()
    _set_paragraph_spacing(spacer, after=1)


def _add_resident_docx_checklist(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    _set_table_layout(table)
    for row_index, pair in enumerate(rows):
        row = table.rows[row_index]
        _set_row_height(row, 0.68, exact=False)
        for col_index, item in enumerate(pair):
            cell = row.cells[col_index]
            cell.width = Cm(8.9)
            _clear_cell(cell)
            _set_cell_border(cell, color='E1E8E2', size='6')
            _set_cell_shading(cell, 'FFFFFF')
            _set_cell_margins(cell, top=70, start=90, bottom=70, end=90)

            paragraph = cell.paragraphs[0]
            _set_paragraph_spacing(paragraph, line=1.0)

            label_run = paragraph.add_run(f'{item[0]}: ')
            label_run.bold = True
            label_run.font.name = 'Arial'
            label_run.font.size = Pt(7.1)
            label_run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x4F)

            value_run = paragraph.add_run(item[1])
            value_run.font.name = 'Arial'
            value_run.font.size = Pt(8.1)
            if item[1] == 'Yes':
                value_run.bold = True
                value_run.font.color.rgb = RGBColor(0x1A, 0x7A, 0x4A)
            else:
                value_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def _add_resident_docx_body_section(document, title, body):
    heading = document.add_paragraph()
    _set_paragraph_spacing(heading, before=2, after=3)
    heading_run = heading.add_run(title)
    heading_run.bold = True
    heading_run.font.name = 'Arial'
    heading_run.font.size = Pt(11)
    heading_run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x4F)

    panel = document.add_table(rows=1, cols=1)
    _set_table_layout(panel)
    cell = panel.cell(0, 0)
    _clear_cell(cell)
    _set_cell_border(cell, color='D7E2DB', size='6')
    _set_cell_shading(cell, 'F8FAF7')
    _set_cell_margins(cell, top=140, start=150, bottom=140, end=150)

    first = cell.paragraphs[0]
    _set_paragraph_spacing(first, line=1.25)
    lines = _split_lines(body)
    for idx, line in enumerate(lines):
        paragraph = first if idx == 0 else cell.add_paragraph()
        _set_paragraph_spacing(paragraph, after=1, line=1.25)
        run = paragraph.add_run(line)
        run.font.name = 'Arial'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x22, 0x35, 0x31)


def generate_resident_pdf(resident):
    try:
        from weasyprint import HTML
    except ImportError:
        return HttpResponse('WeasyPrint is not installed. Cannot generate PDF.', status=500)

    html_string = render_to_string(
        'core/resident_pdf.html',
        {
            'resident': resident,
            'care_home_name': 'Welshwood Manor',
            'resident_photo_uri': _resident_photo_uri(resident),
            'support_needs_text': _clean_text(resident.administration_preferences),
            'mar_guidance_text': _clean_text(resident.mar_front_page_text, fallback=''),
        },
    )
    html = HTML(string=html_string, base_url=str(settings.BASE_DIR))
    pdf = html.write_pdf()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_build_resident_export_filename(resident, "pdf")}"'
    return response



def generate_resident_docx(resident):
    if Document is None:
        return HttpResponse('python-docx is not installed. Cannot generate DOCX.', status=500)

    document = Document()
    _set_document_defaults(document)

    _add_resident_docx_banner(document, 'WELSHWOOD MANOR', fill='E3F0E8', color='1F5C4F', font_size=9, border_color='D3E1D8')
    _add_resident_docx_center_text(document, 'RESIDENT PROFILE', size=8.2, bold=True, color='3A4A4A', after=3)
    _add_resident_docx_photo(document, resident)
    _add_resident_docx_center_text(document, f'BEDROOM NUMBER {resident.room_number or "N/A"}', size=8.6, bold=True, color='4A6A5A', after=3)

    _add_resident_docx_info_table(
        document,
        [
            [('Name', resident.name or 'Not recorded'), ('NOK', resident.next_of_kin or 'Not recorded')],
            [('Date of Birth', _format_date(resident.date_of_birth)), ('NOK Contact', resident.next_of_kin_contact or 'Not recorded')],
            [('Date of Photo', _format_date(resident.updated_at.date() if resident.updated_at else None)), ('NHS Number', resident.nhs_number or 'Not recorded')],
            [('GP Surgery', resident.gp_surgery or resident.gp_name or 'Not recorded'), ('GP Contact', resident.gp_contact or 'Not recorded')],
            [('Residential or Nursing', resident.get_residential_or_nursing_display()), ('DNAR in place', 'YES' if resident.dnar_in_place else 'No')],
            [('Pharmacy', resident.pharmacy_name or 'Not recorded'), ('Pharmacy Contact', resident.pharmacy_contact or 'Not recorded')],
        ],
    )

    if (resident.medical_conditions or '').strip():
        _add_resident_docx_body_section(document, 'MEDICAL CONDITIONS:', resident.medical_conditions)

    if (resident.medication_alerts or '').strip():
        _add_resident_docx_body_section(document, 'MEDICATION ALERTS:', resident.medication_alerts)

    _add_resident_docx_alert(document, f'ALLERGIES: {resident.allergies or "Not recorded"}')

    _add_resident_docx_checklist(
        document,
        [
            [('Is diabetic', 'Yes' if resident.is_diabetic else 'No'), ('Has a PEG in situ', 'Yes' if resident.has_peg else 'No')],
            [('Self medicates', 'Yes' if resident.self_medicates else 'No'), ('Risk assessment in place', 'Yes' if resident.risk_assessment_in_place else 'No')],
            [('Swallowing difficulties', 'Yes' if resident.has_swallowing_difficulties else 'No'), ('On oxygen', 'Yes' if resident.on_oxygen else 'No')],
            [('Can take medication orally', 'Yes' if resident.can_take_medication_orally else 'No'), ('Requires an inhaler', 'Yes' if resident.requires_inhaler else 'No')],
            [('Has capacity around medication', 'Yes' if resident.has_medication_capacity else 'No'), ('Has Parkinsons', 'Yes' if resident.has_parkinsons else 'No')],
            [('MCA in place', 'Yes' if resident.mca_in_place else 'No'), ('Has dementia', 'Yes' if resident.has_dementia else 'No')],
            [('On blood thinning medication', 'Yes' if resident.on_blood_thinning_medication else 'No'), ('Needs pulse or BP before medication', 'Yes' if resident.requires_pulse_or_bp_before_medication else 'No')],
        ],
    )

    document.add_page_break()
    _add_resident_docx_banner(document, 'WELSHWOOD MANOR', fill='E3F0E8', color='1F5C4F', font_size=10, border_color='D3E1D8')
    _add_resident_docx_body_section(
        document,
        f'WHAT LEVEL OF SUPPORT IS NEEDED FOR {resident.name.upper() if resident.name else "THIS RESIDENT"}:',
        resident.administration_preferences or 'Not recorded',
    )
    if (resident.mar_front_page_text or '').strip():
        _add_resident_docx_body_section(document, 'ADDITIONAL MAR GUIDANCE:', resident.mar_front_page_text)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{_build_resident_export_filename(resident, "docx")}"'
    return response
