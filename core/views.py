import json
import logging
import csv
from datetime import datetime
from io import StringIO
from functools import wraps

from django.contrib import messages
from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout
from django.contrib.auth.forms import AuthenticationForm
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.utils.http import url_has_allowed_host_and_scheme
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_GET, require_POST

from .ai_helper import get_ai_resident_profile_assistance
from .forms import ResidentForm, ResidentImportForm
from .models import AuditLog, Resident
from .pdf_generator import (
    generate_resident_docx,
    generate_resident_pdf,
)

logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

DEFAULT_MEDICAL_CONDITION_SUGGESTIONS = [
    'Arthritis',
    'Asthma',
    'COPD',
    'Dementia',
    'Depression',
    'Diabetes',
    'Dysphagia',
    'Epilepsy',
    'Heart failure',
    'Hypertension',
    'Parkinson\'s disease',
    'Stroke',
]

DEFAULT_ALLERGY_SUGGESTIONS = [
    'No known allergies',
    'Penicillin',
    'Latex',
    'Codeine',
    'Sulfa drugs',
    'Aspirin',
    'Ibuprofen',
    'Peanuts',
]

CARE_HOME_NAME = 'Welshwood Manor'


def role_names_for_user(user):
    if not user.is_authenticated:
        return []
    if user.is_superuser:
        return ['Portal Admin']
    return list(user.groups.order_by('name').values_list('name', flat=True))


def has_portal_access(user):
    return user.is_authenticated and (user.is_superuser or user.groups.exists())


def can_view_portal(user):
    return has_portal_access(user) and (user.is_superuser or user.has_perm('core.view_resident'))


def can_edit_portal(user):
    return has_portal_access(user) and (
        user.is_superuser or user.has_perm('core.add_resident') or user.has_perm('core.change_resident')
    )


REQUIRED_PROFILE_FIELDS = {
    'medical_conditions': 'Medical conditions',
    'allergies': 'Allergies',
    'medication_alerts': 'Medication alerts',
    'gp_name': 'GP name',
    'pharmacy_name': 'Pharmacy name',
    'monitoring_requirements': 'Monitoring requirements',
    'administration_preferences': 'Administration preferences',
    'legal_safeguarding_information': 'Legal / safeguarding information',
}


def portal_access_required(check):
    def decorator(view_func):
        @wraps(view_func)
        def _wrapped(request, *args, **kwargs):
            if not request.user.is_authenticated:
                request.session['post_login_next'] = request.get_full_path()
                return redirect('login')
            if not check(request.user):
                raise PermissionDenied('You do not have permission to access this part of the portal.')
            return view_func(request, *args, **kwargs)

        return _wrapped

    return decorator


viewer_required = portal_access_required(can_view_portal)
editor_required = portal_access_required(can_edit_portal)




def _safe_next_url(request, fallback):
    candidate = request.GET.get('next') or request.POST.get('next')
    if candidate and url_has_allowed_host_and_scheme(
        candidate,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return candidate
    return fallback

def log_audit_event(action, description, user=None, target=None, resident=None, metadata=None):
    resident = resident or getattr(target, 'resident', None) or (target if isinstance(target, Resident) else None)
    entity_type = target.__class__.__name__ if target is not None else 'Session'
    entity_id = str(getattr(target, 'pk', '')) if target is not None else ''
    entity_label = str(target) if target is not None else ''
    AuditLog.objects.create(
        user=user if getattr(user, 'is_authenticated', False) else None,
        resident=resident,
        action=action,
        entity_type=entity_type,
        entity_id=entity_id,
        entity_label=entity_label,
        description=description,
        metadata=metadata or {},
    )


@require_GET
def login_view(request):
    if request.user.is_authenticated:
        return redirect('resident_list')
    form = AuthenticationForm(request)
    return render(request, 'registration/login.html', {'form': form})


@require_POST
def login_submit(request):
    form = AuthenticationForm(request, data=request.POST)
    if form.is_valid():
        user = form.get_user()
        if not has_portal_access(user):
            form.add_error(None, 'Your account does not have a portal role assigned yet.')
        else:
            auth_login(request, user)
            log_audit_event('login', 'User logged into the MAR portal.', user=user)
            fallback = request.session.pop('post_login_next', None) or '/'
            return redirect(_safe_next_url(request, fallback))
    return render(request, 'registration/login.html', {'form': form}, status=400)


@viewer_required
@require_POST
def logout_view(request):
    user = request.user
    log_audit_event('logout', 'User logged out of the MAR portal.', user=user)
    auth_logout(request)
    return redirect('login')


def _split_multivalue_suggestions(values):
    suggestions = set()
    for value in values:
        if not value:
            continue
        for line in str(value).replace(';', '\n').splitlines():
            for part in line.split(','):
                cleaned = part.strip()
                if cleaned:
                    suggestions.add(cleaned)
    return suggestions


def get_medical_condition_suggestions():
    raw_values = Resident.objects.order_by().values_list('medical_conditions', flat=True).distinct()
    suggestions = _split_multivalue_suggestions(raw_values)
    suggestions.update(DEFAULT_MEDICAL_CONDITION_SUGGESTIONS)
    return sorted(suggestions, key=str.lower)


def get_allergy_suggestions():
    raw_values = Resident.objects.order_by().values_list('allergies', flat=True).distinct()
    suggestions = _split_multivalue_suggestions(raw_values)
    suggestions.update(DEFAULT_ALLERGY_SUGGESTIONS)
    return sorted(suggestions, key=str.lower)


def get_gp_name_suggestions():
    names = Resident.objects.order_by().values_list('gp_name', flat=True).distinct()
    return sorted({n.strip() for n in names if n and n.strip()}, key=str.lower)


def get_gp_surgery_suggestions():
    surgeries = Resident.objects.order_by().values_list('gp_surgery', flat=True).distinct()
    return sorted({s.strip() for s in surgeries if s and s.strip()}, key=str.lower)


def get_pharmacy_name_suggestions():
    names = Resident.objects.order_by().values_list('pharmacy_name', flat=True).distinct()
    return sorted({n.strip() for n in names if n and n.strip()}, key=str.lower)


def calculate_profile_gaps(data):
    missing = [label for field, label in REQUIRED_PROFILE_FIELDS.items() if not (data.get(field) or '').strip()]
    inconsistencies = []
    allergies = (data.get('allergies') or '').lower()
    medication_alerts = (data.get('medication_alerts') or '').lower()
    medical_conditions = (data.get('medical_conditions') or '').lower()
    monitoring = (data.get('monitoring_requirements') or '').lower()

    allergy_lines = [line.strip() for line in allergies.splitlines() if line.strip()]
    if 'no known allergies' in allergies and len(allergy_lines) > 1:
        inconsistencies.append('Allergy section includes “No known allergies” alongside additional allergy entries.')
    if 'nkda' in medication_alerts and allergy_lines:
        inconsistencies.append('Medication alerts mention NKDA while allergy details are present.')
    if 'diabetes' in medical_conditions and 'blood glucose' not in monitoring and 'blood sugar' not in monitoring:
        inconsistencies.append('Monitoring requirements may be incomplete for diabetes because blood glucose monitoring is not mentioned.')
    if 'warfarin' in medical_conditions and 'bleeding' not in medication_alerts:
        inconsistencies.append('Medication alerts may be incomplete because warfarin-related bleeding precautions are not mentioned.')

    return missing, inconsistencies


def build_dashboard_context(user, residents):
    recent_logs = AuditLog.objects.select_related('user', 'resident')[:8]
    return {
        'residents': residents,
        'total_residents': residents.count(),
        'recent_logs': recent_logs,
        'can_edit': can_edit_portal(user),
        'role_names': role_names_for_user(user),
        'care_home_name': CARE_HOME_NAME,
    }


def build_resident_form_context(request, form, title, resident=None):
    current_data = form.initial.copy()
    if request.method == 'POST':
        current_data.update(request.POST.dict())
    elif resident:
        for field in form.fields:
            current_data[field] = getattr(resident, field, '') if hasattr(resident, field) else form.initial.get(field, '')
    missing_information, inconsistencies = calculate_profile_gaps(current_data)
    return {
        'form': form,
        'title': title,
        'resident': resident,
        'missing_information': missing_information,
        'inconsistencies': inconsistencies,
        'required_profile_fields': REQUIRED_PROFILE_FIELDS,
        'medical_condition_suggestions': get_medical_condition_suggestions(),
        'allergy_suggestions': get_allergy_suggestions(),
        'gp_name_suggestions': get_gp_name_suggestions(),
        'gp_surgery_suggestions': get_gp_surgery_suggestions(),
        'pharmacy_name_suggestions': get_pharmacy_name_suggestions(),
        'care_home_name': CARE_HOME_NAME,
    }


def _csv_value(row, *keys):
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ''


def _parse_csv_date(value, field_label):
    value = (value or '').strip()
    if not value:
        return None

    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'{field_label} must use YYYY-MM-DD or DD/MM/YYYY format.')


def _parse_bool(value):
    normalized = str(value or '').strip().lower()
    return normalized in {'1', 'true', 'yes', 'y', 't'}


def _load_import_rows(upload):
    filename = upload.name.lower()

    if filename.endswith('.csv'):
        try:
            decoded = upload.read().decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('CSV must be UTF-8 encoded.') from exc

        reader = csv.DictReader(StringIO(decoded))
        if not reader.fieldnames:
            raise ValueError('CSV file is empty or missing a header row.')
        return list(reader)

    if filename.endswith('.xlsx'):
        if load_workbook is None:
            raise ValueError('Excel support is not available. Install openpyxl to import .xlsx files.')

        workbook = load_workbook(upload, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError('Excel file is empty or missing a header row.')

        headers = [str(cell).strip() if cell is not None else '' for cell in rows[0]]
        if not any(headers):
            raise ValueError('Excel file is empty or missing a header row.')

        data_rows = []
        for row in rows[1:]:
            row_dict = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else ''
                row_dict[header] = '' if value is None else str(value)
            if any(str(v).strip() for v in row_dict.values()):
                data_rows.append(row_dict)
        return data_rows

    raise ValueError('Please upload a .csv or .xlsx file.')


@editor_required
def resident_import(request):
    form = ResidentImportForm(request.POST or None, request.FILES or None)

    if request.method == 'POST' and form.is_valid():
        upload = form.cleaned_data['csv_file']
        try:
            rows = _load_import_rows(upload)
        except ValueError as exc:
            form.add_error('csv_file', str(exc))
        else:
            created = 0
            updated = 0
            skipped = 0
            row_errors = []

            for index, row in enumerate(rows, start=2):
                try:
                    name = _csv_value(row, 'name', 'resident_name')
                    room_number = _csv_value(row, 'room_number', 'room')
                    dob = _parse_csv_date(_csv_value(row, 'date_of_birth', 'dob'), 'date_of_birth')

                    if not name or not room_number or dob is None:
                        missing = []
                        if not name:
                            missing.append('name')
                        if not room_number:
                            missing.append('room_number')
                        if dob is None:
                            missing.append('date_of_birth')
                        raise ValueError(f'Missing required field(s): {", ".join(missing)}')

                    nhs_number = _csv_value(row, 'nhs_number', 'nhs')
                    residential_or_nursing = _csv_value(row, 'residential_or_nursing', 'care_type') or 'residential'
                    if residential_or_nursing not in {'residential', 'nursing'}:
                        residential_or_nursing = 'residential'

                    defaults = {
                        'name': name,
                        'room_number': room_number,
                        'residential_or_nursing': residential_or_nursing,
                        'date_of_birth': dob,
                        'mental_capacity': _csv_value(row, 'mental_capacity'),
                        'medical_conditions': _csv_value(row, 'medical_conditions'),
                        'allergies': _csv_value(row, 'allergies'),
                        'medication_alerts': _csv_value(row, 'medication_alerts'),
                        'gp_name': _csv_value(row, 'gp_name'),
                        'gp_surgery': _csv_value(row, 'gp_surgery'),
                        'gp_contact': _csv_value(row, 'gp_contact'),
                        'pharmacy_name': _csv_value(row, 'pharmacy_name'),
                        'pharmacy_contact': _csv_value(row, 'pharmacy_contact'),
                        'next_of_kin': _csv_value(row, 'next_of_kin'),
                        'next_of_kin_contact': _csv_value(row, 'next_of_kin_contact'),
                        'monitoring_requirements': _csv_value(row, 'monitoring_requirements'),
                        'administration_preferences': _csv_value(row, 'administration_preferences'),
                        'legal_safeguarding_information': _csv_value(row, 'legal_safeguarding_information'),
                        'mar_front_page_text': _csv_value(row, 'mar_front_page_text'),
                        'dnar_in_place': _parse_bool(_csv_value(row, 'dnar_in_place')),
                        'is_diabetic': _parse_bool(_csv_value(row, 'is_diabetic')),
                        'has_peg': _parse_bool(_csv_value(row, 'has_peg')),
                        'self_medicates': _parse_bool(_csv_value(row, 'self_medicates')),
                        'risk_assessment_in_place': _parse_bool(_csv_value(row, 'risk_assessment_in_place')),
                        'has_swallowing_difficulties': _parse_bool(_csv_value(row, 'has_swallowing_difficulties')),
                        'on_oxygen': _parse_bool(_csv_value(row, 'on_oxygen')),
                        'can_take_medication_orally': _parse_bool(_csv_value(row, 'can_take_medication_orally')),
                        'requires_inhaler': _parse_bool(_csv_value(row, 'requires_inhaler')),
                        'has_medication_capacity': _parse_bool(_csv_value(row, 'has_medication_capacity')),
                        'has_parkinsons': _parse_bool(_csv_value(row, 'has_parkinsons')),
                        'mca_in_place': _parse_bool(_csv_value(row, 'mca_in_place')),
                        'has_dementia': _parse_bool(_csv_value(row, 'has_dementia')),
                        'on_blood_thinning_medication': _parse_bool(_csv_value(row, 'on_blood_thinning_medication')),
                        'requires_pulse_or_bp_before_medication': _parse_bool(_csv_value(row, 'requires_pulse_or_bp_before_medication')),
                    }

                    if nhs_number:
                        resident, was_created = Resident.objects.update_or_create(
                            nhs_number=nhs_number,
                            defaults=defaults | {'nhs_number': nhs_number},
                        )
                    else:
                        resident, was_created = Resident.objects.update_or_create(
                            name=name,
                            date_of_birth=dob,
                            defaults=defaults,
                        )

                    if was_created:
                        created += 1
                    else:
                        updated += 1

                except Exception as exc:
                    skipped += 1
                    row_errors.append(f'Row {index}: {exc}')

            log_audit_event(
                'create',
                'Imported resident data from file.',
                user=request.user,
                metadata={
                    'created': created,
                    'updated': updated,
                    'skipped': skipped,
                    'filename': upload.name,
                },
            )

            if created or updated:
                messages.success(
                    request,
                    f'Import finished. Created {created}, updated {updated}, skipped {skipped}.',
                )
            else:
                messages.warning(request, 'Import completed but no rows were added or updated.')

            for row_error in row_errors[:8]:
                messages.warning(request, row_error)
            if len(row_errors) > 8:
                messages.warning(request, f'{len(row_errors) - 8} additional row errors were omitted from display.')

            return redirect('resident_list')

    return render(
        request,
        'core/resident_import.html',
        {
            'form': form,
            'care_home_name': CARE_HOME_NAME,
        },
    )


@editor_required
@require_GET
def resident_import_template(request):
    headers = [
        'name',
        'room_number',
        'date_of_birth',
        'nhs_number',
        'residential_or_nursing',
        'mental_capacity',
        'medical_conditions',
        'allergies',
        'medication_alerts',
        'gp_name',
        'gp_surgery',
        'gp_contact',
        'pharmacy_name',
        'pharmacy_contact',
        'next_of_kin',
        'next_of_kin_contact',
        'monitoring_requirements',
        'administration_preferences',
        'legal_safeguarding_information',
        'mar_front_page_text',
        'dnar_in_place',
        'is_diabetic',
        'has_peg',
        'self_medicates',
        'risk_assessment_in_place',
        'has_swallowing_difficulties',
        'on_oxygen',
        'can_take_medication_orally',
        'requires_inhaler',
        'has_medication_capacity',
        'has_parkinsons',
        'mca_in_place',
        'has_dementia',
        'on_blood_thinning_medication',
        'requires_pulse_or_bp_before_medication',
    ]

    sample_row = [
        'Jane Smith',
        '12A',
        '1943-05-16',
        '1234567890',
        'residential',
        'Has capacity with occasional prompting.',
        'Dementia\nHypertension',
        'Penicillin',
        'Falls risk - monitor post-medication.',
        'Dr Patel',
        'High Street Surgery',
        '02071234567',
        'Well Pharmacy',
        '02070001111',
        'John Smith',
        '07700123456',
        'Blood pressure weekly',
        'Requires prompting at medication round',
        'Best-interest note filed 2026-01-03',
        'Crushable meds only where clinically approved.',
        'no',
        'no',
        'no',
        'no',
        'no',
        'no',
        'no',
        'yes',
        'no',
        'yes',
        'no',
        'yes',
        'no',
        'no',
        'no',
    ]

    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerow(sample_row)

    response = HttpResponse(buffer.getvalue(), content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="resident_import_template.csv"'
    return response


@viewer_required
def resident_list(request):
    query = request.GET.get('q', '').strip()
    residents = Resident.objects.all()
    if query:
        residents = residents.filter(
            Q(name__icontains=query)
            | Q(room_number__icontains=query)
            | Q(nhs_number__icontains=query)
            | Q(medical_conditions__icontains=query)
        )
    context = build_dashboard_context(request.user, residents)
    context['query'] = query
    return render(request, 'core/resident_list.html', context)


@editor_required
def resident_create(request):
    if request.method == 'POST':
        form = ResidentForm(request.POST, request.FILES)
        if form.is_valid():
            resident = form.save()
            log_audit_event('create', f'Resident profile created for {resident.name}.', user=request.user, target=resident)
            messages.success(request, f'Resident profile for {resident.name} added successfully.')
            return redirect('resident_detail', pk=resident.pk)
    else:
        form = ResidentForm()
    return render(request, 'core/resident_form.html', build_resident_form_context(request, form, 'Create MAR resident profile'))


@viewer_required
def resident_detail(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    missing_information, inconsistencies = calculate_profile_gaps({
        field: getattr(resident, field, '') for field in REQUIRED_PROFILE_FIELDS
    } | {
        'medical_conditions': resident.medical_conditions,
        'allergies': resident.allergies,
        'medication_alerts': resident.medication_alerts,
        'monitoring_requirements': resident.monitoring_requirements,
    })
    context = {
        'resident': resident,
        'missing_information': missing_information,
        'inconsistencies': inconsistencies,
        'audit_logs': resident.audit_logs.select_related('user')[:10],
        'can_edit': can_edit_portal(request.user),
        'care_home_name': CARE_HOME_NAME,
    }
    return render(request, 'core/resident_detail.html', context)


@viewer_required
def resident_profile(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    context = {
        'resident': resident,
        'care_home_name': CARE_HOME_NAME,
        'missing_information': calculate_profile_gaps({field: getattr(resident, field, '') for field in REQUIRED_PROFILE_FIELDS})[0],
    }
    return render(request, 'core/resident_profile.html', context)


@editor_required
def resident_edit(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    if request.method == 'POST':
        form = ResidentForm(request.POST, request.FILES, instance=resident)
        if form.is_valid():
            resident = form.save()
            log_audit_event(
                'update',
                f'Resident profile updated for {resident.name}.',
                user=request.user,
                target=resident,
                metadata={'changed_fields': form.changed_data},
            )
            messages.success(request, f'Resident profile for {resident.name} updated successfully.')
            return redirect('resident_detail', pk=resident.pk)
    else:
        form = ResidentForm(instance=resident)
    return render(request, 'core/resident_form.html', build_resident_form_context(request, form, 'Edit MAR resident profile', resident=resident))


@editor_required
def resident_delete(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    if request.method == 'POST':
        name = resident.name
        log_audit_event('delete', f'Resident profile deleted for {name}.', user=request.user, target=resident)
        resident.delete()
        messages.success(request, f'Resident profile for {name} deleted.')
        return redirect('resident_list')
    return render(request, 'core/resident_confirm_delete.html', {'object': resident, 'type': 'Resident Profile', 'care_home_name': CARE_HOME_NAME})


@viewer_required
def resident_pdf(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    log_audit_event('export_pdf', f'Exported resident MAR profile PDF for {resident.name}.', user=request.user, target=resident)
    return generate_resident_pdf(resident)


@viewer_required
def resident_docx(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    log_audit_event('export_docx', f'Exported resident MAR profile DOCX for {resident.name}.', user=request.user, target=resident)
    return generate_resident_docx(resident)


@viewer_required
def resident_print(request, pk):
    resident = get_object_or_404(Resident, pk=pk)
    log_audit_event('print', f'Opened print-friendly MAR profile for {resident.name}.', user=request.user, target=resident)
    return render(request, 'core/resident_print.html', {'resident': resident, 'care_home_name': CARE_HOME_NAME})


@editor_required
@require_POST
def resident_ai_assist(request):
    try:
        data = json.loads(request.body)
        suggestions = get_ai_resident_profile_assistance(data)
        log_audit_event(
            'ai_assist',
            'Generated AI-assisted resident profile suggestions requiring staff review.',
            user=request.user,
            metadata={'fields_requested': sorted(data.keys())},
        )
        return JsonResponse(suggestions)
    except Exception:
        logger.exception('Error in resident_ai_assist endpoint')
        return JsonResponse({'error': 'An unexpected error occurred. Please try again.'}, status=500)
